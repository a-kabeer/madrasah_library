"""The Excel template, export, and the four-step import.

Covers the workbook mechanics (what a bad file does), the column matching,
what the planner refuses and what it merely warns about, and the guarantee
that matters most: a book group lands whole or not at all.
"""

import io

from django.test import TestCase
from django.urls import reverse

from openpyxl import Workbook, load_workbook

from library import excel as book_excel
from library.models import (
    Author,
    Book,
    BookCopy,
    BookVolume,
    Category,
    Publisher,
)

from .helpers import (
    make_author,
    make_book,
    make_category,
    make_copy,
    make_location,
    make_publisher,
    make_shelf,
    make_user,
    make_volume,
)


LABELS = {column.key: column.label for column in book_excel.COLUMNS}


def workbook_bytes(headings, rows):
    """An .xlsx in memory, the way a user's file arrives."""

    book = Workbook()
    sheet = book.active

    sheet.append(list(headings))

    for row in rows:
        sheet.append(list(row))

    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)

    return buffer


class ImportTestCase(TestCase):

    def setUp(self):
        make_user(username="admin_u", password="pass12345", role="Admin")
        self.client.login(username="admin_u", password="pass12345")

        self.author = make_author(name="Imam Nawawi")
        self.category = make_category(name="Hadith")
        self.publisher = make_publisher(name="Dar Ibn Kathir")

        self.hall = make_location(name="Main Hall")
        self.a1 = make_shelf(location=self.hall, shelf_code="A-1")

        self.annexe = make_location(name="Annexe")
        self.b1 = make_shelf(location=self.annexe, shelf_code="B-1")

        self.url = reverse("book_import")

    def tearDown(self):
        # Each upload writes a file that only a finished or cancelled
        # import removes; a test that stops at the preview would otherwise
        # leave it behind.
        from library.views import import_storage

        try:
            _, names = import_storage.listdir("")

        except OSError:
            return

        for name in names:
            import_storage.delete(name)

    # ------------------------------------------------------------ helpers

    def upload(self, headings, rows, name="books.xlsx"):
        buffer = workbook_bytes(headings, rows)
        buffer.name = name

        return self.client.post(
            self.url, {"action": "upload", "workbook": buffer}
        )

    def mapping_from(self, response):
        return {
            "column_%d" % row["index"]: row["key"]
            for row in response.context["column_rows"]
        }

    def preview(self, headings, rows):
        response = self.upload(headings, rows)
        payload = self.mapping_from(response)
        payload["action"] = "map"

        return self.client.post(self.url, payload)

    def run_import(self, headings, rows):
        preview = self.preview(headings, rows)

        payload = self.mapping_from(preview)
        payload["action"] = "confirm"

        return self.client.post(self.url, payload)

    def full_headings(self, *keys):
        return [LABELS[key] for key in keys]


class TemplateTests(ImportTestCase):

    def test_the_template_downloads_as_a_workbook(self):
        response = self.client.get(reverse("book_import_template"))

        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])

    def test_it_has_the_columns_and_an_instructions_sheet(self):
        response = self.client.get(reverse("book_import_template"))

        book = load_workbook(io.BytesIO(response.content))

        self.assertEqual(book.sheetnames, ["Books", "Instructions"])

        headings = [cell.value for cell in book["Books"][1]]

        self.assertEqual(headings, [c.label for c in book_excel.COLUMNS])

    def test_the_headings_are_readable_not_model_fields(self):
        headings = [c.label for c in book_excel.COLUMNS]

        self.assertIn("Book Title", headings)
        self.assertIn("Number of Copies", headings)

        # No internal field names, and the one id there is says what it is.
        for heading in headings:
            self.assertNotIn("_id", heading)

        self.assertIn("system", LABELS["book_id"])

    def test_it_carries_example_rows_that_show_a_multi_volume_book(self):
        response = self.client.get(reverse("book_import_template"))
        sheet = load_workbook(io.BytesIO(response.content))["Books"]

        titles = [row[1].value for row in sheet.iter_rows(min_row=2, max_row=4)]

        # The last two examples are one book over two rows.
        self.assertEqual(titles[1], titles[2])

    def test_the_books_sheet_holds_only_headings_and_examples(self):
        # A note in column A of that sheet came back as a data row with an
        # unreadable Book ID, so the prose lives on the other sheet.
        response = self.client.get(reverse("book_import_template"))
        sheet = load_workbook(io.BytesIO(response.content))["Books"]

        rows = [
            row for row in sheet.iter_rows(values_only=True)
            if any(str(cell or "").strip() for cell in row)
        ]

        # One heading row plus the three examples, and nothing else.
        self.assertEqual(len(rows), 4)

    def test_the_template_imports_cleanly_once_the_examples_are_removed(self):
        response = self.client.get(reverse("book_import_template"))
        sheet = load_workbook(io.BytesIO(response.content))["Books"]

        headings = [cell.value for cell in sheet[1]]

        # What a librarian actually submits: the template with their own
        # row in place of the examples.
        row = [""] * len(headings)
        row[headings.index(LABELS["title"])] = "From The Template"
        row[headings.index(LABELS["author"])] = "Imam Nawawi"
        row[headings.index(LABELS["copies"])] = "2"
        row[headings.index(LABELS["location"])] = "Main Hall"
        row[headings.index(LABELS["shelf"])] = "A-1"

        result = self.run_import(headings, [row])

        self.assertEqual(result.context["results"]["created"], 1)
        self.assertEqual(result.context["results"]["copies"], 2)
        self.assertEqual(result.context["failures"], [])

    def test_an_assistant_cannot_download_it(self):
        self.client.logout()
        make_user(username="assist", password="pass12345", role="Assistant")
        self.client.login(username="assist", password="pass12345")

        self.assertEqual(
            self.client.get(reverse("book_import_template")).status_code, 403
        )


class UploadTests(ImportTestCase):

    def test_a_good_file_moves_on_to_the_mapping_step(self):
        response = self.upload(
            self.full_headings("title", "author"),
            [["Riyad as-Salihin", "Imam Nawawi"]],
        )

        self.assertEqual(response.context["step"], "map")
        self.assertEqual(response.context["state"]["rows"], 1)

    def test_no_file_is_refused(self):
        response = self.client.post(self.url, {"action": "upload"})

        self.assertEqual(response.context["step"], "upload")
        self.assertIn("Choose an .xlsx", response.context["error"])

    def test_a_csv_renamed_to_xlsx_is_refused(self):
        # The extension is not evidence; the zip signature is.
        fake = io.BytesIO(b"Title,Author\nSomething,Someone\n")
        fake.name = "books.xlsx"

        response = self.client.post(
            self.url, {"action": "upload", "workbook": fake}
        )

        self.assertIn("does not look like an .xlsx", response.context["error"])

    def test_an_empty_file_is_refused(self):
        empty = io.BytesIO(b"")
        empty.name = "books.xlsx"

        response = self.client.post(
            self.url, {"action": "upload", "workbook": empty}
        )

        self.assertIn("empty", response.context["error"])

    def test_a_damaged_workbook_is_refused_without_a_traceback(self):
        broken = io.BytesIO(b"PK\x03\x04 and then nonsense")
        broken.name = "books.xlsx"

        response = self.client.post(
            self.url, {"action": "upload", "workbook": broken}
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("could not be opened", response.context["error"])
        self.assertNotContains(response, "Traceback")
        self.assertNotContains(response, "openpyxl")

    def test_a_sheet_with_only_headings_is_refused(self):
        response = self.upload(self.full_headings("title", "author"), [])

        self.assertIn("no rows underneath", response.context["error"])

    def test_a_sheet_with_no_headings_is_refused(self):
        response = self.upload([], [])

        self.assertIn("no column headings", response.context["error"])

    def test_blank_rows_are_skipped_not_imported(self):
        response = self.upload(
            self.full_headings("title", "author"),
            [
                ["Riyad as-Salihin", "Imam Nawawi"],
                ["", ""],
                ["Al-Adhkar", "Imam Nawawi"],
            ],
        )

        self.assertEqual(response.context["state"]["rows"], 2)

    def test_extra_columns_are_kept_but_ignored(self):
        response = self.upload(
            self.full_headings("title", "author") + ["Shelf Note", "Price"],
            [["Riyad as-Salihin", "Imam Nawawi", "n/a", "100"]],
        )

        keys = [row["key"] for row in response.context["column_rows"]]

        self.assertEqual(keys, ["title", "author", "", ""])

    def test_nothing_is_written_by_uploading(self):
        before = Book.objects.count()

        self.upload(
            self.full_headings("title", "author"),
            [["Riyad as-Salihin", "Imam Nawawi"]],
        )

        self.assertEqual(Book.objects.count(), before)


class MappingTests(ImportTestCase):

    def test_exact_headings_are_matched(self):
        response = self.upload(
            self.full_headings("title", "author", "copies"),
            [["A", "B", "1"]],
        )

        self.assertEqual(
            [r["key"] for r in response.context["column_rows"]],
            ["title", "author", "copies"],
        )

    def test_aliases_and_odd_casing_are_matched(self):
        # Nobody should have to rename their spreadsheet first.
        response = self.upload(
            ["BOOK NAME", "writer", "Qty", "vol no", "Shelf Code"],
            [["A", "B", "2", "1", "A-1"]],
        )

        self.assertEqual(
            [r["key"] for r in response.context["column_rows"]],
            ["title", "author", "copies", "volume_number", "shelf"],
        )

    def test_a_heading_matched_twice_is_only_used_once(self):
        response = self.upload(
            ["Book Title", "Title"],
            [["A", "B"]],
        )

        self.assertEqual(
            [r["key"] for r in response.context["column_rows"]],
            ["title", ""],
        )

    def test_a_column_can_be_ignored_by_hand(self):
        response = self.upload(
            self.full_headings("title", "author", "category"),
            [["A", "B", "Hadith"]],
        )

        payload = self.mapping_from(response)
        payload["column_2"] = book_excel.IGNORE
        payload["action"] = "map"

        preview = self.client.post(self.url, payload)

        self.assertEqual(preview.context["step"], "preview")

        group = preview.context["groups"][0]
        self.assertEqual(group.category, "")

    def test_a_missing_required_mapping_is_refused(self):
        response = self.upload(
            self.full_headings("title", "author"),
            [["A", "B"]],
        )

        payload = self.mapping_from(response)
        payload["column_1"] = book_excel.IGNORE
        payload["action"] = "map"

        back = self.client.post(self.url, payload)

        self.assertEqual(back.context["step"], "map")
        self.assertIn("Author", back.context["error"])
        self.assertIn("required", back.context["error"])

    def test_the_same_field_mapped_twice_is_refused(self):
        self.upload(["One", "Two", "Three"], [["A", "B", "C"]])

        payload = {
            "action": "map",
            "column_0": "title",
            "column_1": "author",
            "column_2": "title",
        }

        back = self.client.post(self.url, payload)

        self.assertEqual(back.context["step"], "map")
        self.assertIn("more than one column", back.context["error"])

    def test_a_file_with_no_recognisable_headings_still_maps_by_hand(self):
        response = self.upload(["Col A", "Col B"], [["A", "B"]])

        self.assertEqual(
            [r["key"] for r in response.context["column_rows"]], ["", ""]
        )

        preview = self.client.post(self.url, {
            "action": "map", "column_0": "title", "column_1": "author",
        })

        self.assertEqual(preview.context["step"], "preview")


class PreviewTests(ImportTestCase):

    def test_the_preview_says_what_will_happen_and_writes_nothing(self):
        before = Book.objects.count()

        response = self.preview(
            self.full_headings("title", "author", "copies", "location", "shelf"),
            [["Riyad as-Salihin", "Imam Nawawi", "3", "Main Hall", "A-1"]],
        )

        summary = response.context["summary"]

        self.assertEqual(response.context["step"], "preview")
        self.assertEqual(summary["new_books"], 1)
        self.assertEqual(summary["new_copies"], 3)
        self.assertEqual(summary["skipped"], 0)
        self.assertEqual(Book.objects.count(), before)

    def test_a_new_related_record_is_flagged_as_a_warning_not_an_error(self):
        response = self.preview(
            self.full_headings("title", "author", "category"),
            [["Something", "A Brand New Author", "A Brand New Subject"]],
        )

        group = response.context["groups"][0]

        self.assertFalse(group.blocked)
        self.assertTrue(
            any("will be created" in p.message for p in group.problems)
        )

    def test_an_existing_title_and_author_warns_rather_than_blocks(self):
        # Nothing in the schema makes a title unique, so this cannot be an
        # error — a library may hold two editions.
        make_book(title="Riyad as-Salihin", author=self.author)

        response = self.preview(
            self.full_headings("title", "author"),
            [["Riyad as-Salihin", "Imam Nawawi"]],
        )

        group = response.context["groups"][0]

        self.assertFalse(group.blocked)
        self.assertTrue(
            any("already has" in p.message for p in group.problems)
        )

    def test_a_missing_title_blocks_the_group(self):
        response = self.preview(
            self.full_headings("title", "author"),
            [["", "Imam Nawawi"]],
        )

        self.assertTrue(response.context["groups"][0].blocked)
        self.assertEqual(response.context["summary"]["skipped"], 1)

    def test_an_unknown_location_blocks_the_row(self):
        response = self.preview(
            self.full_headings("title", "author", "copies", "location", "shelf"),
            [["A", "Imam Nawawi", "1", "Nowhere", "A-1"]],
        )

        group = response.context["groups"][0]

        self.assertTrue(group.blocked)
        self.assertTrue(
            any("No location called" in p.message for p in group.problems)
        )

    def test_a_shelf_from_the_wrong_location_blocks_the_row(self):
        # B-1 exists, but in Annexe, not Main Hall.
        response = self.preview(
            self.full_headings("title", "author", "copies", "location", "shelf"),
            [["A", "Imam Nawawi", "1", "Main Hall", "B-1"]],
        )

        group = response.context["groups"][0]

        self.assertTrue(group.blocked)
        self.assertTrue(
            any("has no shelf" in p.message for p in group.problems)
        )

    def test_copies_without_a_place_block_the_row(self):
        response = self.preview(
            self.full_headings("title", "author", "copies"),
            [["A", "Imam Nawawi", "2"]],
        )

        self.assertTrue(response.context["groups"][0].blocked)

    def test_a_place_without_copies_is_only_a_warning(self):
        response = self.preview(
            self.full_headings("title", "author", "location", "shelf"),
            [["A", "Imam Nawawi", "Main Hall", "A-1"]],
        )

        group = response.context["groups"][0]

        self.assertFalse(group.blocked)
        self.assertTrue(
            any("no copies" in p.message for p in group.problems)
        )


class AddBooksTests(ImportTestCase):

    def test_a_plain_book_with_no_copies(self):
        self.run_import(
            self.full_headings("title", "author"),
            [["Al-Adhkar", "Imam Nawawi"]],
        )

        book = Book.objects.get(title="Al-Adhkar")

        self.assertEqual(book.author_id, self.author.id)
        self.assertEqual(BookCopy.objects.count(), 0)

        # No volume named and nothing to place, so nothing is invented. The
        # implicit volume exists to hang a copy off; a blank cell is not
        # the Add Book form's explicit "Single volume" choice.
        self.assertEqual(BookVolume.objects.filter(book=book).count(), 0)

    def test_a_named_volume_is_created_even_with_no_copies(self):
        self.run_import(
            self.full_headings("title", "author", "volume_number"),
            [["Named", "Imam Nawawi", "2"]],
        )

        volumes = BookVolume.objects.filter(book__title="Named")

        self.assertEqual(volumes.count(), 1)
        self.assertEqual(volumes.first().volume_number, 2)

    def test_a_single_volume_book_with_copies(self):
        self.run_import(
            self.full_headings("title", "author", "copies", "location", "shelf"),
            [["Al-Adhkar", "Imam Nawawi", "3", "Main Hall", "A-1"]],
        )

        book = Book.objects.get(title="Al-Adhkar")
        volumes = BookVolume.objects.filter(book=book)

        # One implicit volume, unnamed, exactly as the Add Book form makes.
        self.assertEqual(volumes.count(), 1)
        self.assertEqual(volumes.first().volume_number, 1)
        self.assertEqual(volumes.first().title, "")

        copies = BookCopy.objects.filter(volume=volumes.first())
        self.assertEqual(copies.count(), 3)
        self.assertEqual({c.shelf_id for c in copies}, {self.a1.id})

    def test_a_multi_volume_book_from_several_rows(self):
        headings = self.full_headings(
            "title", "author", "volume_number", "volume_title",
            "copies", "location", "shelf",
        )

        self.run_import(headings, [
            ["Sahih al-Bukhari", "Imam Bukhari", "1", "Kitab al-Iman",
             "2", "Main Hall", "A-1"],
            ["Sahih al-Bukhari", "Imam Bukhari", "2", "Kitab al-Salah",
             "1", "Main Hall", "A-1"],
            ["Sahih al-Bukhari", "Imam Bukhari", "3", "", "0", "", ""],
        ])

        # One book, not three.
        self.assertEqual(Book.objects.filter(title="Sahih al-Bukhari").count(), 1)

        book = Book.objects.get(title="Sahih al-Bukhari")
        volumes = BookVolume.objects.filter(book=book).order_by("volume_number")

        self.assertEqual(
            [(v.volume_number, v.title) for v in volumes],
            [(1, "Kitab al-Iman"), (2, "Kitab al-Salah"), (3, "")],
        )

        self.assertEqual(
            [BookCopy.objects.filter(volume=v).count() for v in volumes],
            [2, 1, 0],
        )

    def test_two_books_in_one_file(self):
        self.run_import(
            self.full_headings("title", "author"),
            [
                ["Al-Adhkar", "Imam Nawawi"],
                ["Riyad as-Salihin", "Imam Nawawi"],
            ],
        )

        self.assertEqual(Book.objects.count(), 2)

    def test_one_volume_can_be_spread_over_two_shelves(self):
        # Real libraries do this, and the export writes a row per shelf, so
        # a repeated volume number is a second placement, not a clash.
        headings = self.full_headings(
            "title", "author", "volume_number", "copies", "location", "shelf"
        )
        b2 = make_shelf(location=self.hall, shelf_code="A-2")

        response = self.run_import(headings, [
            ["Split", "Imam Nawawi", "1", "2", "Main Hall", "A-1"],
            ["Split", "Imam Nawawi", "1", "1", "Main Hall", "A-2"],
        ])

        book = Book.objects.get(title="Split")
        volumes = BookVolume.objects.filter(book=book)

        # One volume, three copies, on the two shelves the rows named.
        self.assertEqual(volumes.count(), 1)
        self.assertEqual(response.context["results"]["volumes"], 1)

        copies = BookCopy.objects.filter(volume=volumes.first())

        self.assertEqual(copies.count(), 3)
        self.assertEqual(copies.filter(shelf=self.a1).count(), 2)
        self.assertEqual(copies.filter(shelf=b2).count(), 1)

    def test_a_repeated_volume_number_is_reported_as_a_warning(self):
        response = self.preview(
            self.full_headings("title", "author", "volume_number"),
            [
                ["Sahih al-Bukhari", "Imam Bukhari", "1"],
                ["Sahih al-Bukhari", "Imam Bukhari", "1"],
            ],
        )

        group = response.context["groups"][0]

        self.assertFalse(group.blocked)
        self.assertTrue(
            any("also on row" in p.message and not p.blocking
                for p in group.problems)
        )

    def test_related_records_are_created_and_reused(self):
        self.run_import(
            self.full_headings("title", "author", "category", "publisher"),
            [
                ["One", "New Author", "New Subject", "New Press"],
                ["Two", "New Author", "New Subject", "New Press"],
            ],
        )

        # Created once each, not once per row.
        self.assertEqual(Author.objects.filter(name="New Author").count(), 1)
        self.assertEqual(Category.objects.filter(name="New Subject").count(), 1)
        self.assertEqual(Publisher.objects.filter(name="New Press").count(), 1)

    def test_an_existing_related_name_is_matched_case_insensitively(self):
        self.run_import(
            self.full_headings("title", "author", "category"),
            [["One", "imam nawawi", "HADITH"]],
        )

        book = Book.objects.get(title="One")

        self.assertEqual(book.author_id, self.author.id)
        self.assertEqual(book.category_id, self.category.id)
        self.assertEqual(Author.objects.count(), 1)

    def test_the_import_is_logged_against_the_user(self):
        from library.models import ActivityLog

        self.run_import(
            self.full_headings("title", "author"),
            [["One", "Imam Nawawi"]],
        )

        log = ActivityLog.objects.filter(action="IMPORT").first()

        self.assertIsNotNone(log)
        self.assertEqual(log.user.username, "admin_u")
        self.assertIn("1 book(s) added", log.description)


class CopyCodeTests(ImportTestCase):

    def headings(self):
        return self.full_headings(
            "title", "author", "copies", "copy_codes", "location", "shelf"
        )

    def test_codes_are_generated_when_none_are_given(self):
        self.run_import(
            self.full_headings("title", "author", "copies", "location", "shelf"),
            [["One", "Imam Nawawi", "2", "Main Hall", "A-1"]],
        )

        codes = list(BookCopy.objects.values_list("copy_code", flat=True))

        self.assertEqual(len(codes), 2)

        # From the one generator the Add Book form uses.
        for code in codes:
            self.assertTrue(code.startswith("LIB-"), code)

        for copy in BookCopy.objects.all():
            self.assertEqual(copy.copy_code, "LIB-%06d" % copy.id)

    def test_manual_codes_are_used_as_given(self):
        self.run_import(self.headings(), [
            ["One", "Imam Nawawi", "2", "ACC-1, ACC-2", "Main Hall", "A-1"],
        ])

        self.assertEqual(
            sorted(BookCopy.objects.values_list("copy_code", flat=True)),
            ["ACC-1", "ACC-2"],
        )

    def test_the_wrong_number_of_manual_codes_blocks_the_row(self):
        response = self.preview(self.headings(), [
            ["One", "Imam Nawawi", "3", "ACC-1, ACC-2", "Main Hall", "A-1"],
        ])

        group = response.context["groups"][0]

        self.assertTrue(group.blocked)
        self.assertTrue(
            any("one per copy" in p.message for p in group.problems)
        )

    def test_a_code_repeated_inside_the_file_blocks_it(self):
        response = self.preview(self.headings(), [
            ["One", "Imam Nawawi", "1", "SAME", "Main Hall", "A-1"],
            ["Two", "Imam Nawawi", "1", "SAME", "Main Hall", "A-1"],
        ])

        blocked = [g for g in response.context["groups"] if g.blocked]

        self.assertEqual(len(blocked), 1)
        self.assertTrue(
            any("twice in this file" in p.message
                for g in blocked for p in g.problems)
        )

    def test_a_code_already_in_the_library_blocks_the_row(self):
        make_copy(shelf=self.a1, copy_code="TAKEN-1")

        response = self.preview(self.headings(), [
            ["One", "Imam Nawawi", "1", "TAKEN-1", "Main Hall", "A-1"],
        ])

        group = response.context["groups"][0]

        self.assertTrue(group.blocked)
        self.assertTrue(
            any("already in use" in p.message for p in group.problems)
        )

    def test_a_code_clashing_only_in_case_blocks_the_row(self):
        make_copy(shelf=self.a1, copy_code="Taken-2")

        response = self.preview(self.headings(), [
            ["One", "Imam Nawawi", "1", "TAKEN-2", "Main Hall", "A-1"],
        ])

        self.assertTrue(response.context["groups"][0].blocked)

    def test_codes_without_copies_block_the_row(self):
        response = self.preview(self.headings(), [
            ["One", "Imam Nawawi", "", "ACC-9", "Main Hall", "A-1"],
        ])

        self.assertTrue(response.context["groups"][0].blocked)


class UpdateTests(ImportTestCase):
    """Export, edit, re-import — landing on the same record."""

    def setUp(self):
        super().setUp()

        self.book = make_book(
            title="Original Title",
            author=self.author,
            category=self.category,
        )
        self.volume = make_volume(
            book=self.book, volume_number=1, title=""
        )
        self.copy = make_copy(
            volume=self.volume, shelf=self.a1, copy_code="LIB-990001"
        )

    def export_rows(self):
        response = self.client.get(reverse("book_export"))
        sheet = load_workbook(io.BytesIO(response.content))["Books"]

        headings = [cell.value for cell in sheet[1]]
        rows = [
            ["" if cell.value is None else cell.value for cell in row]
            for row in sheet.iter_rows(min_row=2)
        ]

        return headings, rows

    def test_the_export_downloads_and_uses_the_template_columns(self):
        response = self.client.get(reverse("book_export"))

        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])

        headings, rows = self.export_rows()

        self.assertEqual(headings, [c.label for c in book_excel.COLUMNS])
        self.assertEqual(len(rows), 1)

    def test_the_export_carries_the_identifier_and_the_data(self):
        headings, rows = self.export_rows()

        row = dict(zip(headings, rows[0]))

        self.assertEqual(row[LABELS["book_id"]], self.book.id)
        self.assertEqual(row[LABELS["title"]], "Original Title")
        self.assertEqual(row[LABELS["author"]], self.author.name)
        self.assertEqual(row[LABELS["copies"]], 1)
        self.assertEqual(row[LABELS["copy_codes"]], "LIB-990001")
        self.assertEqual(row[LABELS["location"]], "Main Hall")
        self.assertEqual(row[LABELS["shelf"]], "A-1")

        # A single implicit volume is not called "Volume 1" here either.
        self.assertEqual(row[LABELS["volume_number"]], "")

    def test_a_round_trip_updates_rather_than_duplicates(self):
        headings, rows = self.export_rows()

        # Left exactly as exported apart from the title — which is the
        # ordinary way this is used, and the case that has to be clean.
        rows[0][headings.index(LABELS["title"])] = "Corrected Title"

        self.run_import(headings, rows)

        self.assertEqual(Book.objects.count(), 1)

        self.book.refresh_from_db()
        self.assertEqual(self.book.title, "Corrected Title")

    def test_a_round_trip_leaves_the_inventory_alone(self):
        headings, rows = self.export_rows()
        rows[0][headings.index(LABELS["title"])] = "Renamed"

        self.run_import(headings, rows)

        self.volume.refresh_from_db()
        self.copy.refresh_from_db()

        self.assertEqual(BookVolume.objects.filter(book=self.book).count(), 1)
        self.assertEqual(self.copy.copy_code, "LIB-990001")
        self.assertEqual(self.copy.shelf_id, self.a1.id)

        # And no second copy was made from the code the export mentioned.
        self.assertEqual(BookCopy.objects.count(), 1)

    def test_an_unchanged_export_re_imports_with_nothing_to_do(self):
        # Including a volume whose copies sit on two shelves, which the
        # export writes as two rows.
        second = make_shelf(location=self.hall, shelf_code="A-2")
        make_copy(
            volume=self.volume, shelf=second, copy_code="LIB-990002"
        )

        headings, rows = self.export_rows()

        response = self.run_import(headings, rows)
        results = response.context["results"]

        self.assertEqual(results["created"], 0)
        self.assertEqual(results["updated"], 1)
        self.assertEqual(results["copies"], 0)
        self.assertEqual(results["volumes"], 0)
        self.assertEqual(results["failed"], [])
        self.assertEqual(response.context["failures"], [])

        # Nothing added, nothing lost.
        self.assertEqual(BookCopy.objects.count(), 2)
        self.assertEqual(BookVolume.objects.filter(book=self.book).count(), 1)

    def test_a_code_belonging_to_another_book_is_still_refused(self):
        other = make_book(title="Elsewhere", author=self.author)
        other_volume = make_volume(book=other, volume_number=1, title="")
        make_copy(
            volume=other_volume, shelf=self.a1, copy_code="LIB-880001"
        )

        headings, rows = self.export_rows()

        # The export is sorted by title, so the row has to be found rather
        # than assumed — "Elsewhere" comes before "Original Title".
        title_at = headings.index(LABELS["title"])
        mine = next(
            row for row in rows if row[title_at] == "Original Title"
        )
        mine[headings.index(LABELS["copy_codes"])] = "LIB-880001"

        response = self.preview(headings, rows)

        group = next(
            g for g in response.context["groups"]
            if g.title == "Original Title"
        )

        self.assertTrue(group.blocked)
        self.assertTrue(
            any("already in use" in p.message for p in group.problems)
        )

    def test_a_new_code_on_an_update_row_does_add_a_copy(self):
        headings, rows = self.export_rows()

        rows[0][headings.index(LABELS["copies"])] = 2
        rows[0][headings.index(LABELS["copy_codes"])] = "LIB-990001, NEW-1"

        self.run_import(headings, rows)

        self.assertEqual(BookCopy.objects.count(), 2)
        self.assertTrue(BookCopy.objects.filter(copy_code="NEW-1").exists())

    def test_a_bibliographic_only_round_trip_creates_nothing(self):
        # 203 empty volumes were once created by exactly this: books with
        # no volumes exporting as a row with a blank volume number.
        plain = make_book(title="No Inventory", author=self.author)

        headings, rows = self.export_rows()

        response = self.preview(headings, rows)
        summary = response.context["summary"]

        self.assertEqual(summary["new_volumes"], 0)

        payload = self.mapping_from(response)
        payload["action"] = "confirm"
        done = self.client.post(self.url, payload)

        self.assertEqual(done.context["results"]["volumes"], 0)
        self.assertEqual(BookVolume.objects.filter(book=plain).count(), 0)

    def test_the_preview_promises_the_volumes_it_will_create(self):
        headings, rows = self.export_rows()

        # Ask for a volume the book does not have.
        rows[0][headings.index(LABELS["volume_number"])] = 5

        response = self.preview(headings, rows)

        self.assertEqual(response.context["summary"]["new_volumes"], 1)

        payload = self.mapping_from(response)
        payload["action"] = "confirm"
        done = self.client.post(self.url, payload)

        self.assertEqual(done.context["results"]["volumes"], 1)

    def test_the_preview_calls_it_an_update(self):
        headings, rows = self.export_rows()

        response = self.preview(headings, rows)

        group = response.context["groups"][0]

        self.assertEqual(group.action, "update")
        self.assertEqual(response.context["summary"]["updated_books"], 1)
        self.assertEqual(response.context["summary"]["new_books"], 0)

    def test_a_blank_optional_cell_does_not_clear_the_field(self):
        headings, rows = self.export_rows()
        rows[0][headings.index(LABELS["category"])] = ""

        self.run_import(headings, rows)

        self.book.refresh_from_db()

        self.assertEqual(self.book.category_id, self.category.id)

    def test_an_unknown_book_id_blocks_the_group(self):
        response = self.preview(
            self.full_headings("book_id", "title", "author"),
            [["999999", "Ghost", "Imam Nawawi"]],
        )

        group = response.context["groups"][0]

        self.assertTrue(group.blocked)
        self.assertTrue(any("No book has ID" in p.message
                            for p in group.problems))

    def test_a_non_numeric_book_id_blocks_the_group(self):
        response = self.preview(
            self.full_headings("book_id", "title", "author"),
            [["abc", "Ghost", "Imam Nawawi"]],
        )

        self.assertTrue(response.context["groups"][0].blocked)

    def test_an_export_can_be_narrowed_the_way_the_list_is(self):
        make_book(title="Somewhere Else", author=make_author(name="Other"))

        response = self.client.get(
            reverse("book_export"), {"search": "Original"}
        )
        sheet = load_workbook(io.BytesIO(response.content))["Books"]

        titles = [row[1].value for row in sheet.iter_rows(min_row=2)]

        self.assertEqual(titles, ["Original Title"])

    def test_an_assistant_cannot_export(self):
        self.client.logout()
        make_user(username="assist", password="pass12345", role="Assistant")
        self.client.login(username="assist", password="pass12345")

        self.assertEqual(
            self.client.get(reverse("book_export")).status_code, 403
        )


class SafetyTests(ImportTestCase):

    def test_a_good_book_still_imports_when_another_is_bad(self):
        headings = self.full_headings(
            "title", "author", "copies", "location", "shelf"
        )

        self.run_import(headings, [
            ["Good One", "Imam Nawawi", "1", "Main Hall", "A-1"],
            ["Bad One", "Imam Nawawi", "1", "Nowhere At All", "A-1"],
            ["Good Two", "Imam Nawawi", "1", "Main Hall", "A-1"],
        ])

        titles = sorted(Book.objects.values_list("title", flat=True))

        self.assertEqual(titles, ["Good One", "Good Two"])
        self.assertEqual(BookCopy.objects.count(), 2)

    def test_a_blocked_book_leaves_no_volume_or_copy_behind(self):
        headings = self.full_headings(
            "title", "author", "volume_number", "copies", "location", "shelf"
        )

        # Volume 1 is fine; volume 2 names a shelf in the wrong location.
        self.run_import(headings, [
            ["Half Bad", "Imam Nawawi", "1", "1", "Main Hall", "A-1"],
            ["Half Bad", "Imam Nawawi", "2", "1", "Main Hall", "B-1"],
        ])

        self.assertFalse(Book.objects.filter(title="Half Bad").exists())
        self.assertEqual(BookVolume.objects.count(), 0)
        self.assertEqual(BookCopy.objects.count(), 0)

    def test_the_results_page_reports_both_sides(self):
        headings = self.full_headings(
            "title", "author", "copies", "location", "shelf"
        )

        response = self.run_import(headings, [
            ["Good One", "Imam Nawawi", "2", "Main Hall", "A-1"],
            ["Bad One", "Imam Nawawi", "1", "Nowhere", "A-1"],
        ])

        results = response.context["results"]

        self.assertEqual(response.context["step"], "done")
        self.assertEqual(results["created"], 1)
        self.assertEqual(results["copies"], 2)
        self.assertEqual(len(response.context["failures"]), 1)

    def test_an_error_report_can_be_downloaded_afterwards(self):
        self.run_import(
            self.full_headings("title", "author", "copies", "location", "shelf"),
            [["Bad One", "Imam Nawawi", "1", "Nowhere", "A-1"]],
        )

        response = self.client.get(reverse("book_import_errors"))

        self.assertEqual(response.status_code, 200)

        sheet = load_workbook(io.BytesIO(response.content))["Errors"]
        rows = [
            [cell.value for cell in row]
            for row in sheet.iter_rows(min_row=2)
        ]

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 2)          # the spreadsheet row
        self.assertEqual(rows[0][1], "Bad One")
        self.assertIn("No location called", rows[0][2])

    def test_cancelling_forgets_the_upload(self):
        self.upload(
            self.full_headings("title", "author"),
            [["One", "Imam Nawawi"]],
        )

        response = self.client.post(self.url, {"action": "cancel"})

        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            self.client.get(self.url).context["step"], "upload"
        )

    def test_confirming_without_an_upload_is_harmless(self):
        response = self.client.post(self.url, {"action": "confirm"})

        self.assertEqual(response.context["step"], "upload")
        self.assertIn("expired", response.context["error"])
        self.assertEqual(Book.objects.count(), 0)

    def test_an_assistant_cannot_reach_the_import(self):
        self.client.logout()
        make_user(username="assist", password="pass12345", role="Assistant")
        self.client.login(username="assist", password="pass12345")

        self.assertEqual(self.client.get(self.url).status_code, 403)

        response = self.upload(
            self.full_headings("title", "author"),
            [["Sneaky", "Someone"]],
        )

        self.assertEqual(response.status_code, 403)
        self.assertEqual(Book.objects.count(), 0)


class PerformanceTests(ImportTestCase):

    def test_the_plan_does_not_query_per_row(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        headings = self.full_headings(
            "title", "author", "copies", "location", "shelf"
        )

        def rows(count):
            return [
                ["Book %d" % index, "Imam Nawawi", "1", "Main Hall", "A-1"]
                for index in range(count)
            ]

        with CaptureQueriesContext(connection) as few:
            self.preview(headings, rows(2))

        with CaptureQueriesContext(connection) as many:
            self.preview(headings, rows(40))

        # The lookups are fetched once up front; the planner works from
        # dictionaries after that.
        self.assertEqual(len(few), len(many))
